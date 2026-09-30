"""Read-only source extraction. Cached cells; no formula results invented.
BoE workbook obtained by B2; row coordinates preserved in slim xlsx.
"""
from pathlib import Path
import openpyxl,csv,hashlib,json
ROOT=Path(__file__).resolve().parents[4]
OUT=Path(__file__).resolve().parent
SRC=ROOT/'downloads/B2_BoE_millennium_v31.xlsx'
w=openpyxl.load_workbook(SRC,read_only=True,data_only=True)
slim=openpyxl.Workbook();slim.remove(slim.active)
records=[]
for name in ['A1. Headline series','A4. Ind Production 1270-1870','A40. Trade by region 1710-1822','A41. Trade by region 1784+','A47. Wages and prices']:
 s=w[name];d=slim.create_sheet(name[:31])
 for rid,row in enumerate(s,1):
  first=row[0].value
  selected=(rid<=8 or (isinstance(first,(int,float)) and 1795<=first<=1850) or (name.startswith('A41') and rid<=150))
  if selected:
   for c in row:
    if c.value is not None:
     d.cell(c.row,c.column,c.value)
     records.append([name,c.coordinate,c.value])
slim.save(OUT/'B5_boe_selected.xlsx')
with open(OUT/'B5_boe_cells.csv','w') as f:
 writer=csv.writer(f);writer.writerow(['source_sheet','source_cell','cached_value']);writer.writerows(records)
with open(OUT/'B5_boe_manifest.json','w') as f:
 json.dump({'source':str(SRC.relative_to(ROOT)),'sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'method':'openpyxl read_only data_only; source coordinates preserved in selected workbook','warning':'cached source formulas not recalculated; mixed reconstructed and historical series'},f,indent=2)
print('saved',len(records),'cells; sheets',slim.sheetnames)
