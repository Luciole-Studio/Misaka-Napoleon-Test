"""B2 retained-data extraction. No re-estimation of BoE sources.
Source workbook v3.1 is too large for the Office reader; copy selected source
sheets as cached-value evidence with original row/column coordinates preserved.
Do not overwrite or recalculate the downloaded workbook.
"""
from pathlib import Path
import openpyxl, json, hashlib
ROOT=Path(__file__).resolve().parents[4]
OUT=Path(__file__).resolve().parent
SRC=ROOT/'downloads/B2_BoE_millennium_v31.xlsx'
SHEETS=['Disclaimer','Front page','Corrections to V3.1','A1. Headline series','A9. Nominal GDP (A)','Notes on GDP estimates','A27. Central govt borrowing ','A29. The National Debt','A30a. Nat Debt mkt vals 1727','A31. Interest rates & asset ps ','M10. Mthly long-term rates']
src=openpyxl.load_workbook(SRC,read_only=True,data_only=True)
out=openpyxl.Workbook(); out.remove(out.active)
for name in SHEETS:
    target=out.create_sheet(name)
    for row in src[name].iter_rows():
        for cell in row:
            if cell.value is not None: target.cell(cell.row,cell.column,cell.value)
    print(name,src[name].max_row,src[name].max_column)
out.save(OUT/'BoE_selected_cached_values.xlsx')
record={'source_url':'https://www.bankofengland.co.uk/-/media/boe/files/statistics/research-datasets/a-millennium-of-macroeconomic-data-for-the-uk.xlsx','source':str(SRC.relative_to(ROOT)), 'sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'vintage':'BoE millennium v3.1, updated through 2016; downloaded 2026-09-23','retrieved_at':'2026-09-23','units':'GBP million (fiscal); percent (yield/ratios); as specified in original sheet headers','frequency':'annual, fiscal/calendar and date conventions as source headers; monthly rate sheet separately','seasonal_adjustment':'NSA historical annual observations','real_or_nominal':'nominal current GBP and nominal yields; historical GDP estimates explicitly identified','base_year':'not applicable to nominal GBP and ratios; no deflation in B2 fiscal series','extraction':'openpyxl read_only=True data_only=True; cached source values, original coordinates; selected sheets only; not source formula recalculation','sheets':SHEETS}
(OUT/'BoE_provenance.json').write_text(json.dumps(record,ensure_ascii=False,indent=2))
