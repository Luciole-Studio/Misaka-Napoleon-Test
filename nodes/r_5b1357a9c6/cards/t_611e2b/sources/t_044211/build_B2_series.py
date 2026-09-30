"""Extract source-defined calendar estimates from the retained BoE workbook.
No interpolation beyond the BoE source's own date conversions. Amounts GBP million.
Interest/revenue is a computed ratio, NOT a contemporaneous published statistic.
The pre-1801 revenue boundary is GB; national debt is consolidated GB+Ireland
plus the BoE capital estimate of terminable annuities. Keep these distinct.
"""
from pathlib import Path
import openpyxl,csv,json,hashlib
P=Path(__file__).resolve().parent
w=openpyxl.load_workbook(P/'BoE_selected_cached_values.xlsx',read_only=True,data_only=True)
f=w['A27. Central govt borrowing ']; d=w['A29. The National Debt']; y=w['A31. Interest rates & asset ps ']; m=w['M10. Mthly long-term rates']
fields=['year','fiscal_source_label','fiscal_end_convention','revenue_boundary','revenue_gbp_m','expenditure_gbp_m','interest_gbp_m','interest_revenue_pct','primary_expenditure_gbp_m','primary_surplus_gbp_m','cash_deficit_gbp_m','debt_funded_unfunded_calendar_gbp_m','capital_value_terminable_annuities_gbp_m','debt_total_calendar_gbp_m','debt_gdp_pct_boe_all_ireland','consol_yield_annual_pct','direct_income_tax_category_fy_gbp_m','fiscal_source_cells','debt_source_cells','yield_source_cell','income_tax_note']
rows=[]
for year in range(1793,1849):
 fr=year-1688+12;dr=year-1688+6;yr=year-1688+8
 R=f[f'AW{fr}'].value;E=f[f'AT{fr}'].value;I=f[f'AV{fr}'].value
 assert f[f'AR{fr}'].value==year and d[f'AN{dr}'].value==year and y[f'A{yr}'].value==year
 debt=d[f'AO{dr}'].value; face=d[f'Y{dr}'].value; ann=d[f'AB{dr}'].value
 assert face is not None and ann is not None  # all 56 source rows retain both components
 assert abs(R-E+I-f[f'AZ{fr}'].value)<1e-7
 assert abs(debt-face-(ann or 0))<1e-6
 a=[year,f[f'A{fr}'].value,f[f'B{fr}'].value,'GB before 1801; UK from 1801',R,E,I,100*I/R,E-I,R-E+I,E-R,face,ann,debt,d[f'AP{dr}'].value,y[f'T{yr}'].value,f[f'AA{fr}'].value,f'A27!AT{fr};AV{fr};AW{fr};AZ{fr}',f'A29!Y{dr};AB{dr};AO{dr};AP{dr}',f'A31!T{yr}','FY source classification; do not equate enactment-year assessment and cash receipts; 1799 row is bridge quarter; 1815/16 means ending Jan 1816']
 record=dict(zip(fields,a))
 record['debt_component_note']='Source separately reports debt and capitalized terminable annuities throughout 1793-1848; total is not purely funded debt'
 rows.append(record)
# Sparse separately dated series: never add stocks, assessments and fiscal flows.
bank=openpyxl.load_workbook(P.parents[3]/'downloads/B2_BoE_balance_sheet.xlsx',read_only=True,data_only=True)["A1. Bank of England B'Sheet"]
extra=['boe_notes_end_feb_gbp_m','boe_coin_bullion_end_feb_gbp_m','bank_source_cells','property_tax_net_produce_year_to_april5_gbp_m','property_tax_source','new_loan_cash_coupon_cost_pct','new_loan_cost_scope_source']
fields+=extra+['debt_component_note']
pt={1804:.363877,1805:3.919108,1806:4.481958,1807:7.000032,1808:10.817595,1809:11.279423}
loan={1810:4+4/20+2/240,1811:355937.5/7500000*100,1815:5.62}
for r in rows:
 year=r['year'];br=year-1791+101
 if year<=1844:
  assert bank[f'A{br}'].value==year
  r['boe_notes_end_feb_gbp_m']=bank[f'S{br}'].value
  r['boe_coin_bullion_end_feb_gbp_m']=bank[f'O{br}'].value
  r['bank_source_cells']=f'BoE annual balance sheet A1!S{br};O{br}; end-Feb; not annual means'
 r['property_tax_net_produce_year_to_april5_gbp_m']=pt.get(year,'')
 r['property_tax_source']='Hansard 20 June 1809 vol14 Finance Resolutions no15; Apr5 year end' if year in pt else ''
 r['new_loan_cash_coupon_cost_pct']=loan.get(year,'')
 r['new_loan_cost_scope_source']=({1810:'Hansard 20 May 1811 c216 retrospective comparator: 4l4s2d per 100 cash; original 1810 stock package not reconstructed; not complete IRR',1811:'Hansard 20 May 1811 cc214-216; 3/4pct mixed stock plus long annuity; annual coupon/committed cash; excludes payment-timing discount and market-value bonus; not IRR',1815:'Hansard 14 June 1815 cc801-804; June loan coupon package per cash subscribed; not all-year issuance average; excludes payment-timing discount and bonus'}.get(year,''))
with (P/'B2_series.csv').open('w',newline='') as out:
 cw=csv.DictWriter(out,fieldnames=fields);cw.writeheader();cw.writerows(rows)
# Independent monthly yield window, exact monthly source not an imputed smooth path.
mr=[]
for cells in m.iter_rows(min_row=7,values_only=True):
 if isinstance(cells[0],(int,float)) and 1809<=cells[0]<=1812:
  mr.append([int(cells[0]),cells[1],cells[2],cells[10],'BoE M10 C / K, Neal (1990), month end'])
with (P/'B2_consol_monthly_1809_1812.csv').open('w',newline='') as out:
 cw=csv.writer(out);cw.writerow(['year','month','raw_consol_yield_pct','spliced_consol_yield_pct','source']);cw.writerows(mr)
print('rows',len(rows),'monthly',len(mr))
for r in rows:
 if r['year'] in [1793,1797,1803,1806,1810,1811,1812,1813,1815,1816,1822,1830,1848]:
  print(r['year'],*[round(r[k],3) for k in ['revenue_gbp_m','interest_gbp_m','interest_revenue_pct','cash_deficit_gbp_m','debt_total_calendar_gbp_m','consol_yield_annual_pct']])
